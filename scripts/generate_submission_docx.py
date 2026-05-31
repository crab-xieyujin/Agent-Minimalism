from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
OUT_PATH = OUT_DIR / "Agent-Minimalism-作品说明材料.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_doc_defaults(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Title", 22, RGBColor(28, 48, 74)),
        ("Heading 1", 15, RGBColor(28, 79, 138)),
        ("Heading 2", 12.5, RGBColor(38, 38, 38)),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True


def add_paragraph(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "D9EAF7")
        set_cell_text(cell, header, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)

    document.add_paragraph()
    return table


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    set_doc_defaults(document)

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Agent Minimalism 作品说明材料")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("参赛作品：Agent Minimalism | 作品类型：Codex Skill / Agent 工作流设计审查工具")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(90, 90, 90)

    document.add_heading("一、作品概述", level=1)
    add_paragraph(
        document,
        "Agent Minimalism 是一个用于设计、审查和优化 Agent 工作流的 Codex Skill。它的核心原则是：默认工作流化，只把真正存在不确定性的环节交给 Agent。作品面向 AI 产品团队、开发者、自动化流程设计者和企业 AI 落地团队，帮助他们把复杂任务拆分为规则、单次 LLM、固定工作流、局部 Agent 和动态 Agent 五个层级，从而减少 token 消耗、降低延迟、提升稳定性和可维护性。",
    )

    add_table(
        document,
        ["项目", "说明"],
        [
            ["作品名称", "Agent Minimalism"],
            ["一句话描述", "帮团队把复杂 Agent 流程拆成规则、LLM、工作流与局部 Agent，降低成本并提升稳定性。"],
            ["适配产品", "CodeBuddy、WorkBuddy"],
            ["应用行业", "专业服务、软件开发、企业 AI 自动化、AI 产品设计"],
            ["应用场景", "Agent 工作流评审、AI 自动化架构设计、token 成本优化、失败恢复设计、流程工程化改造"],
        ],
    )

    document.add_heading("二、使用场景", level=1)
    add_paragraph(
        document,
        "该 Skill 适用于任何需要设计或改造 AI Agent 系统的场景，尤其适合黑客松、企业 AI 应用搭建、研发提效工具、内容生产流水线、销售运营自动化、客户支持自动化等任务。它不是替代业务 Agent，而是作为一个“架构审查助手”，帮助团队判断哪些地方应该使用 Agent，哪些地方应该降级为更稳定、更便宜的实现方式。",
    )
    add_bullets(
        document,
        [
            "黑客松项目：在短时间内检查方案是否过度 Agent 化，帮助团队快速收敛为可演示、可落地的架构。",
            "AI 产品设计：在 PRD 或技术方案阶段，对 Agent 节点、工具调用、上下文传递方式进行审查。",
            "企业自动化流程：把原本全靠 Agent 执行的流程拆解成规则、模板、脚本、固定 DAG 和少量 Agent 节点。",
            "成本与稳定性优化：定位 token 膨胀、共享上下文过大、步骤不可控、失败无法恢复等问题。",
            "研发与运维场景：用于代码修复、CI 诊断、文档生成、发布流水线等复杂流程的 Agent 边界设计。",
        ],
    )

    document.add_heading("三、解决的问题", level=1)
    add_paragraph(
        document,
        "很多 Agent 系统的问题并不是模型能力不足，而是架构过度 Agent 化：把确定性的格式转换、表单填写、数据校验、固定步骤执行都交给 Agent，导致成本高、速度慢、结果不稳定、调试困难。Agent Minimalism 解决的是“什么时候该用 Agent，什么时候不该用 Agent”的设计判断问题。",
    )
    add_bullets(
        document,
        [
            "降低 token 成本：避免多个 Agent 共享完整上下文，减少重复传递大段历史和工具结果。",
            "降低延迟：把确定性步骤从模型循环中移出，用规则、脚本或固定工作流执行。",
            "提升稳定性：让主路径由可验证、可测试的 L0-L2 步骤组成，只在不确定点使用 Agent。",
            "提升可维护性：明确每个 Agent 节点的职责、工具、停止条件和失败回退方式。",
            "减少架构复杂度：避免为了“看起来智能”而引入不必要的 Planner、多 Agent 协作和动态路由。",
        ],
    )

    document.add_heading("四、核心思路", level=1)
    add_paragraph(
        document,
        "作品采用“复杂度路由”的方法，把一个任务链中的每个步骤归类到最低可行复杂度层级。只要低层级能够可靠完成，就不升级到更复杂的 Agent 方案。Agent 的价值被限定在不确定性、探索、恢复和开放式判断中，而不是包办所有执行环节。",
    )
    add_table(
        document,
        ["层级", "适用情况", "推荐实现"],
        [
            ["L0 Rule", "输入、检查、转换或输出是确定性的", "代码、配置、Schema 校验、正则、模板"],
            ["L1 Single LLM", "需要语义处理，但不需要循环和工具观察", "一次 LLM 调用，最好带结构化输出"],
            ["L2 Fixed Workflow", "步骤较多，但路径基本确定", "Pipeline、DAG、状态机"],
            ["L3 Local Agent", "单个节点不确定，需要工具、重试或判断", "只在该节点内使用 Agent"],
            ["L4 Dynamic Agent", "目标开放，路径无法预先确定", "Planner、工具、记忆、迭代执行"],
        ],
    )

    document.add_heading("五、输入设计", level=1)
    add_paragraph(
        document,
        "用户可以输入自然语言描述、现有工作流、Agent 方案、自动化流程、工具调用链、产品需求或失败案例。Skill 会从输入中识别任务目标、输入输出、步骤链路、决策点、工具依赖、失败模式和上下文传递方式。",
    )
    add_table(
        document,
        ["输入类型", "示例"],
        [
            ["任务目标", "帮我设计一个客服工单自动处理 Agent。"],
            ["现有流程", "用户提交表单 -> Agent 读取需求 -> Agent 生成报价 -> Agent 发邮件。"],
            ["失败问题", "当前多 Agent 流程 token 很高，结果不稳定，调试困难。"],
            ["约束条件", "希望延迟低、成本可控、结果可验证，尽量少用动态 Agent。"],
            ["业务产物", "PRD、流程图、工具列表、API 文档、运营 SOP、研发流水线描述。"],
        ],
    )

    document.add_heading("六、输出设计", level=1)
    add_paragraph(
        document,
        "Skill 的输出不是泛泛建议，而是一份可执行的架构审查结果。它会明确当前方案的问题、每一步的最低复杂度分类、推荐架构、每个 Agent 节点的使用理由，以及 token 和稳定性影响。",
    )
    add_bullets(
        document,
        [
            "当前工作流诊断：指出主路径、确定性步骤、不确定性步骤、失败风险和 token 风险。",
            "步骤分类表：为每个步骤标注 L0-L4，并给出推荐实现方式。",
            "推荐架构：说明哪些步骤应该规则化、模板化、工作流化，哪些地方保留 Agent。",
            "Agent 节点说明：逐一说明每个 Agent 处理的不确定性、可用工具、停止条件和失败回退。",
            "成本与稳定性影响：定性说明 token、延迟、上下文共享、稳定性和可维护性的变化。",
            "实施步骤：给出下一步如何改造或落地的操作建议。",
        ],
    )

    document.add_heading("七、边界与不适用场景", level=1)
    add_paragraph(
        document,
        "Agent Minimalism 是一个设计审查与架构建议 Skill，不是通用执行 Agent，也不直接替用户完成所有业务自动化。它关注的是“如何把流程设计得更可靠”，而不是替代具体业务系统、数据库、API 或前端应用。",
    )
    add_bullets(
        document,
        [
            "不负责直接执行完整业务流程，例如真正发送邮件、提交表单、发布内容或修改生产数据。",
            "不替代领域专家判断，例如法律、医疗、财务等高风险领域的最终决策。",
            "不保证所有任务都应减少 Agent；当目标开放、路径未知、需要探索时，L4 动态 Agent 仍然合理。",
            "不处理没有明确目标或缺少基本输入输出描述的任务；这种情况需要先补充任务定义。",
            "不把成本最低作为唯一目标；在需要创造性、探索性或异常恢复能力时，会保留必要的 Agent 节点。",
        ],
    )

    document.add_heading("八、示例", level=1)
    document.add_heading("示例输入", level=2)
    add_paragraph(
        document,
        "我们想做一个公众号文章生产 Agent：输入热点新闻，Agent 搜索资料、筛选角度、写文章、生成封面、排版、发布到草稿箱。现在流程很慢，而且每一步都把完整上下文传给下一个 Agent。请帮我优化。",
    )
    document.add_heading("示例输出摘要", level=2)
    add_table(
        document,
        ["步骤", "推荐层级", "说明"],
        [
            ["热点新闻结构化", "L0 Rule", "固定字段抽取与校验可用 Schema 完成。"],
            ["选题角度生成", "L1 Single LLM", "语义判断明确，一次结构化生成即可。"],
            ["资料搜索与可信度判断", "L3 Local Agent", "需要观察网页结果并做判断，保留局部 Agent。"],
            ["文章写作", "L1 Single LLM / L2 Fixed Workflow", "可拆为固定写作模板加一次或少量 LLM 调用。"],
            ["封面生成", "L2 Fixed Workflow", "提示词模板、图片生成、尺寸校验可固定化。"],
            ["草稿箱发布", "L2 Fixed Workflow", "登录、上传、填表、保存是固定步骤，失败时再调用恢复 Agent。"],
        ],
    )

    document.add_heading("九、作品价值", level=1)
    add_paragraph(
        document,
        "Agent Minimalism 的价值在于把 Agent 设计从“堆能力”转向“做架构”。它帮助团队在早期就识别哪些智能是必要的，哪些智能其实应该被工程化替代。对于黑客松作品，它可以作为一个面向开发者和企业团队的实用工具，帮助参赛者、产品经理和工程师快速设计出更轻、更稳、更容易落地的 AI 应用。",
    )

    document.add_heading("十、仓库与交付物", level=1)
    add_table(
        document,
        ["交付物", "路径或说明"],
        [
            ["GitHub 仓库", "https://github.com/crab-xieyujin/Agent-Minimalism"],
            ["Skill 主文件", "SKILL.md"],
            ["审查清单", "references/review-checklist.md"],
            ["复杂度路由模板", "references/router-template.md"],
            ["作品说明材料", "docs/Agent-Minimalism-作品说明材料.docx"],
        ],
    )

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Agent Minimalism - 作品说明材料")
    footer_run.font.name = "Microsoft YaHei"
    footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(120, 120, 120)

    document.save(OUT_PATH)


if __name__ == "__main__":
    build_document()
    print(OUT_PATH)
